import os
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from pypdf import PdfReader

try:
    import docx
except ImportError:
    docx = None

from src.models import ParsedResume

# Standard skill dictionary with normalized display names
SKILL_PATTERNS: Dict[str, str] = {
    # AI / Agentic / RAG
    r"\blanggraph\b": "LangGraph",
    r"\blangchain\b": "LangChain",
    r"\bllamaindex\b|\bllama-index\b": "LlamaIndex",
    r"\bcrewai\b|\bcrew-ai\b": "CrewAI",
    r"\bautogen\b": "AutoGen",
    r"\bgoogle adk\b|\badk\b": "Google ADK",
    r"\brag\b|\bretrieval[- ]augmented\b": "RAG",
    r"\bvector (?:db|database|search)\b|\bvectorstore\b": "Vector Search",
    r"\bchroma\b|\bchromadb\b": "ChromaDB",
    r"\bpinecone\b": "Pinecone",
    r"\bfaiss\b": "FAISS",
    r"\bweaviate\b": "Weaviate",
    r"\bqdrant\b": "Qdrant",
    r"\bembeddings?\b": "Embeddings",
    r"\btool[- ]calling\b|\bfunction[- ]calling\b": "Tool Calling",
    r"\bmulti[- ]agent\b|\bagentic\b": "Multi-Agent Systems",
    r"\bllms?\b|\blarge language models?\b": "LLM",
    r"\bopenai\b": "OpenAI",
    r"\bgemini\b": "Google Gemini",
    r"\bclaude\b|\banthropic\b": "Anthropic Claude",
    r"\bhugging\s*face\b|\btransformers\b": "Hugging Face",
    r"\bfine[- ]tuning\b": "Fine-Tuning",
    r"\bollama\b": "Ollama",
    r"\bvllm\b": "vLLM",
    r"\bdeepseek\b": "DeepSeek",
    r"\bnvidia nim\b": "Nvidia NIM",
    r"\bprompt engineering\b": "Prompt Engineering",

    # Python & Backend
    r"\bpython\b|\bpython3\b": "Python",
    r"\bfastapi\b": "FastAPI",
    r"\bflask\b": "Flask",
    r"\bdjango\b": "Django",
    r"\basyncio\b|\basynchronous programming\b": "Asyncio",
    r"\bpostgresql\b|\bpostgres\b": "PostgreSQL",
    r"\bmysql\b": "MySQL",
    r"\bsqlite\b": "SQLite",
    r"\bmongodb\b": "MongoDB",
    r"\bredis\b": "Redis",
    r"\bsqlalchemy\b": "SQLAlchemy",
    r"\bcelery\b": "Celery",
    r"\brest\s*apis?\b|\brestful\b": "REST APIs",
    r"\bgraphql\b": "GraphQL",
    r"\brabbitmq\b": "RabbitMQ",
    r"\bkafka\b": "Kafka",

    # Cloud, DevOps & Full Stack
    r"\bdocker\b": "Docker",
    r"\bkubernetes\b|\bk8s\b": "Kubernetes",
    r"\bgcp\b|\bgoogle cloud\b": "GCP",
    r"\baws\b|\bamazon web services\b": "AWS",
    r"\bazure\b": "Azure",
    r"\blinux\b": "Linux",
    r"\bgit\b|\bgithub\b": "Git",
    r"\breact(?:\.js)?\b": "React",
    r"\bnext(?:\.js)?\b": "Next.js",
    r"\btypescript\b": "TypeScript",
    r"\bjavascript\b": "JavaScript",
    r"\bnode(?:\.js)?\b": "Node.js",
    r"\bexpress(?:\.js)?\b": "Express.js",
    r"\btailwind\b": "Tailwind CSS",
    r"\bstreamlit\b": "Streamlit",

    # Engineering Depth
    r"\bpytest\b": "Pytest",
    r"\bunit tests?\b|\bunittest\b|\btdd\b": "Unit Testing",
    r"\bci/cd\b|\bgithub actions\b": "CI/CD",
    r"\bobservability\b|\bopentelemetry\b": "Observability",
    r"\bprometheus\b": "Prometheus",
    r"\bgrafana\b": "Grafana",
    r"\brate limit(?:ing)?\b": "Rate Limiting",
    r"\bcaching\b": "Caching",
    r"\bconcurrency\b|\bmultithreading\b": "Concurrency",
}

SECTION_HEADERS = {
    "projects": [r"\bprojects?\b", r"\bacademic projects?\b", r"\bkey projects?\b", r"\btechnical projects?\b"],
    "experience": [r"\bexperience\b", r"\bwork experience\b", r"\binternships?\b", r"\bemployment\b", r"\bprofessional experience\b"],
    "skills": [r"\bskills\b", r"\btechnical skills\b", r"\btechnologies\b", r"\bcore competencies\b", r"\bproficiencies\b"],
    "education": [r"\beducation\b", r"\bacademics?\b", r"\bqualifications?\b"],
    "summary": [r"\bsummary\b", r"\bprofessional summary\b", r"\bprofile\b", r"\babout me\b"],
}

class ResumeParser:
    def __init__(self):
        pass

    def extract_text_and_links_pdf(self, file_path: str) -> Tuple[str, List[str]]:
        """Extract text and hyperlinked URLs from a PDF."""
        text_parts = []
        links = []
        reader = PdfReader(file_path)
        for page in reader.pages:
            try:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            except Exception:
                pass

            # Extract annotation hyperlinks
            if "/Annots" in page:
                try:
                    annots = page["/Annots"]
                    for annot in annots:
                        obj = annot.get_object()
                        if "/A" in obj and "/URI" in obj["/A"]:
                            uri = obj["/A"]["/URI"]
                            if isinstance(uri, str):
                                links.append(uri.strip())
                except Exception:
                    pass

        raw_text = "\n".join(text_parts)
        return raw_text, links

    def extract_text_docx(self, file_path: str) -> str:
        """Extract text from a DOCX file."""
        if not docx:
            raise ImportError("python-docx is not installed")
        doc = docx.Document(file_path)
        text_parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text.strip())
        return "\n".join(text_parts)

    def extract_text_txt(self, file_path: str) -> str:
        """Extract text from a TXT file."""
        for enc in ("utf-8", "latin-1", "cp1252"):
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except Exception:
                continue
        return ""

    def extract_email(self, text: str) -> Optional[str]:
        match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
        return match.group(0) if match else None

    def extract_phone(self, text: str) -> Optional[str]:
        match = re.search(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3,5}\)?[-.\s]?\d{3,5}[-.\s]?\d{3,5}", text)
        return match.group(0).strip() if match else None

    def extract_github(self, text: str, links: List[str]) -> Tuple[Optional[str], Optional[str]]:
        """Extract GitHub URL and GitHub username from links and text."""
        # Check explicit links from annotations first
        for link in links:
            m = re.search(r"github\.com/([a-zA-Z0-9_-]+)", link, re.IGNORECASE)
            if m:
                username = m.group(1).strip()
                if username.lower() not in {"join", "login", "features", "explore", "pricing"}:
                    return f"https://github.com/{username}", username

        # Fallback to regex in raw text
        m = re.search(r"(?:https?://(?:www\.)?github\.com/|github\.com/|github\s*:\s*)([a-zA-Z0-9_-]+)", text, re.IGNORECASE)
        if m:
            username = m.group(1).strip()
            if username.lower() not in {"join", "login", "features", "explore", "pricing"}:
                return f"https://github.com/{username}", username

        return None, None

    def extract_linkedin(self, text: str, links: List[str]) -> Optional[str]:
        for link in links:
            if "linkedin.com/in/" in link.lower():
                return link
        m = re.search(r"https?://(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+", text, re.IGNORECASE)
        return m.group(0) if m else None

    def extract_candidate_name(self, text: str, file_name: str, email: Optional[str]) -> str:
        """Heuristically and reliably extract candidate name."""
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        noise_words = {
            "resume", "curriculum vitae", "cv", "page", "summary", "profile",
            "contact", "phone", "email", "education", "experience", "projects", "1", "2", "/"
        }

        role_words = {
            "engineer", "developer", "intern", "associate", "analyst", "architect",
            "lead", "specialist", "scientist", "manager", "fresher", "student", "consultant"
        }

        # Check top 6 lines
        for line in lines[:6]:
            clean_line = re.sub(r"[\u2022\u25cf\u25cb\u2642\u2640\u2191\u2193\t]", " ", line).strip()
            # If line starts with noise word or punctuation
            if not clean_line or clean_line.lower() in noise_words:
                continue
            if clean_line.startswith(("/", "•", "-", "*", "1", "2")):
                clean_line = clean_line.lstrip("/•-* 1234567890").strip()
            if not clean_line:
                continue

            # Skip lines that are job titles / summaries
            lower_clean = clean_line.lower()
            if any(rw in lower_clean for rw in role_words):
                continue
            if "with passion" in lower_clean or "looking for" in lower_clean or "aspiring" in lower_clean:
                continue

            # Remove email / phone / location if concatenated
            clean_line = re.sub(r"Email:.*", "", clean_line, flags=re.I).strip()
            clean_line = re.sub(r"Phone:.*", "", clean_line, flags=re.I).strip()
            clean_line = re.sub(r"Mobile:.*", "", clean_line, flags=re.I).strip()
            clean_line = re.sub(r"Dehradun.*", "", clean_line, flags=re.I).strip()
            clean_line = re.sub(r"\+?\d[\d -]{7,}", "", clean_line).strip()
            clean_line = re.sub(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-.]+", "", clean_line).strip()

            words = clean_line.split()
            # A valid name is usually 1 to 4 alphabetical words
            if 1 <= len(words) <= 4:
                if all(re.match(r"^[A-Za-z.'-]+$", w) for w in words):
                    if clean_line.lower() not in noise_words and len(clean_line) > 2:
                        return clean_line.title()

        # Fallback 1: Derive from email
        if email:
            prefix = email.split("@")[0]
            # Strip trailing numbers
            clean_prefix = re.sub(r"\d+$", "", prefix)
            parts = re.split(r"[._-]", clean_prefix)
            parts = [p.capitalize() for p in parts if p.isalpha() and len(p) > 1]
            if parts:
                return " ".join(parts)

        # Fallback 2: From filename
        base = Path(file_name).stem.replace("_", " ").replace("-", " ")
        return base.title()

    def extract_skills(self, text: str) -> List[str]:
        """Match technologies against defined dictionary using boundary matching."""
        found = []
        for pattern, skill_name in SKILL_PATTERNS.items():
            if re.search(pattern, text, re.IGNORECASE):
                if skill_name not in found:
                    found.append(skill_name)
        return found

    def segment_sections(self, text: str) -> Dict[str, str]:
        """Split resume into common logical sections."""
        lines = text.split("\n")
        current_section = "header"
        sections: Dict[str, List[str]] = {
            "header": [],
            "summary": [],
            "skills": [],
            "projects": [],
            "experience": [],
            "education": [],
        }

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            # Check if line matches a known section header
            matched = False
            for sec_name, patterns in SECTION_HEADERS.items():
                if any(re.match(f"^{p}:?$", line_str, re.IGNORECASE) for p in patterns):
                    current_section = sec_name
                    matched = True
                    break
            
            if not matched:
                sections.setdefault(current_section, []).append(line_str)

        return {k: "\n".join(v) for k, v in sections.items() if v}

    def parse_file(self, file_path: str) -> ParsedResume:
        """Parse resume file from disk into a ParsedResume object."""
        path = Path(file_path)
        ext = path.suffix.lower()
        links: List[str] = []

        if ext == ".pdf":
            raw_text, links = self.extract_text_and_links_pdf(file_path)
        elif ext == ".docx":
            raw_text = self.extract_text_docx(file_path)
        elif ext == ".txt":
            raw_text = self.extract_text_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

        if not raw_text.strip():
            raise ValueError("File is empty or contains no readable text")

        email = self.extract_email(raw_text)
        phone = self.extract_phone(raw_text)
        github_url, github_username = self.extract_github(raw_text, links)
        linkedin_url = self.extract_linkedin(raw_text, links)
        candidate_name = self.extract_candidate_name(raw_text, path.name, email)
        skills = self.extract_skills(raw_text)
        sections = self.segment_sections(raw_text)

        return ParsedResume(
            file_path=str(path.resolve()),
            file_name=path.name,
            raw_text=raw_text,
            candidate_name=candidate_name,
            email=email,
            phone=phone,
            github_url=github_url,
            github_username=github_username,
            linkedin_url=linkedin_url,
            skills=skills,
            sections=sections,
        )
